import asyncio
import copy
import json
import os
import tempfile
import subprocess
import sys
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

from openpyxl import load_workbook

import src.config as config_module
import src.database as database_module
from src.browser import BrowserManager
from src.config import ConfigManager
from src.llm.providers import LLMSettings, LMStudioProvider
from src.models import JobRecord
from src.scraper import LinkedInScraper


class PersistentHistoryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        folder = Path(self.temp.name)
        self.patches = patch.multiple(database_module, DB_FILE=folder / "jobs.db",
                                      EXPORTS_DIR=folder / "exports")
        self.patches.start()
        self.addCleanup(self.temp.cleanup)
        self.addCleanup(self.patches.stop)
        self.db = database_module.DatabaseManager()

    def job(self, number=1, status="REVIEWED", **kwargs):
        return JobRecord(job_url=f"https://www.linkedin.com/jobs/view/{number}/",
                         title="Analyst", company="Fixture Company", status=status, **kwargs)

    def test_review_states_never_enter_applied_history_or_excel(self):
        for index, state in enumerate(("REVIEWED", "READY_FOR_REVIEW", "FAILED")):
            self.db.save_job_record(self.job(index, state))
        self.assertEqual(self.db.get_applied_jobs(), [])
        reports = self.db.export_job_reports()
        workbook = load_workbook(reports["applied"])
        self.assertEqual(workbook.active.max_row, 1)
        workbook.close()
        with self.assertRaises(ValueError):
            self.db.log_applied_job("Analyst", "Fixture", "", "", "https://example.test/job", status="REVIEWED")

    def test_confirmed_external_submission_is_atomic_and_survives_restart(self):
        job = self.job(status="APPLIED", application_url="https://jobs.lever.co/fixture/1",
                       ats_type="Lever", date_applied="2026-09-08T10:00:00-04:00")
        self.db.save_job_record(job)
        self.db.save_job_record(job)
        self.db.save_job_record(self.job(status="FAILED"))
        restarted = database_module.DatabaseManager()
        restarted.recover_interrupted_applications()
        self.assertEqual(len(restarted.get_applied_jobs()), 1)
        self.assertEqual(restarted.get_reviewed_jobs()[0]["status"], "APPLIED")
        reports = restarted.export_job_reports()
        for path in database_module.EXPORTS_DIR.glob("*.xlsx"):
            path.unlink()
        reports = restarted.export_job_reports()
        workbook = load_workbook(reports["applied"])
        headers = {c.value: c.column for c in workbook.active[1]}
        self.assertEqual(workbook.active.cell(2, headers["ATS"]).value, "Lever")
        self.assertEqual(workbook.active.cell(2, headers["Status"]).value, "APPLIED")
        workbook.close()

    def test_second_reader_does_not_expire_review_but_startup_does(self):
        job = self.job(status="READY_FOR_REVIEW")
        self.db.save_job_record(job)
        self.db.queue_application_review(job.title, job.company, job.job_url)
        reader = database_module.DatabaseManager()
        self.assertEqual(len(reader.get_application_reviews()), 1)
        reader.recover_interrupted_applications()
        self.assertEqual(reader.get_application_reviews(), [])
        self.assertEqual(reader.get_application_reviews(None)[0]["status"], "expired")
        self.assertFalse(reader.is_job_known(job.job_url))
        self.assertEqual(reader.get_applied_jobs(), [])

    def test_new_process_can_read_history_and_regenerate_excel(self):
        self.db.save_job_record(self.job(status="APPLIED", date_applied="2026-09-08T10:00:00-04:00"))
        env = dict(os.environ, JOB_APPLIER_DATA_DIR=self.temp.name,
                   JOB_APPLIER_EXPORTS_DIR=str(database_module.EXPORTS_DIR))
        result = subprocess.run([sys.executable, "-c",
            "from src.database import db_manager; "
            "assert len(db_manager.get_applied_jobs()) == 1; "
            "assert all(p.exists() for p in db_manager.export_job_reports().values())"],
            cwd=Path(__file__).resolve().parents[1], env=env, capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_legacy_confirmed_external_row_is_repaired_on_restart(self):
        self.db.save_job_record(self.job(status="APPLIED", date_applied="2026-09-08T10:00:00-04:00"))
        with self.db.get_connection() as conn:
            conn.execute("DELETE FROM applied_jobs")
            conn.commit()
        restarted = database_module.DatabaseManager()
        self.assertEqual(len(restarted.get_applied_jobs()), 1)

    def test_review_decisions_cannot_skip_approval_or_change_terminal_outcome(self):
        job = self.job()
        review = self.db.queue_application_review(job.title, job.company, job.job_url)
        self.assertFalse(self.db.resolve_application_review(review, "submitted"))
        self.assertTrue(self.db.resolve_application_review(review, "approved"))
        self.assertTrue(self.db.resolve_application_review(review, "submitted"))
        self.assertFalse(self.db.resolve_application_review(review, "rejected"))

    def test_teaching_resolved_question_allows_retry(self):
        job = self.job(status="NEEDS_USER_INPUT")
        self.db.save_job_record(job)
        self.db.log_flagged_job(job.title, job.company, job.job_url, "Missing answer")
        self.assertTrue(self.db.is_job_known(job.job_url))
        self.db.resolve_flagged_job(self.db.get_flagged_jobs()[0]["id"])
        self.assertFalse(self.db.is_job_known(job.job_url))
        self.db.log_flagged_job(job.title, job.company, job.job_url, "New question")
        self.assertEqual(self.db.get_flagged_jobs()[0]["reason_flagged"], "New question")

    def test_tracking_and_slug_variants_do_not_duplicate_records(self):
        first = self.job()
        second = self.job()
        second.job_url = "https://www.linkedin.com/jobs/view/analyst-1/?trackingId=x"
        self.db.save_job_record(first)
        self.db.save_job_record(second)
        self.assertEqual(first.job_id, second.job_id)
        self.assertEqual(len(self.db.get_reviewed_jobs()), 1)

    def test_concurrent_history_writers_and_exports(self):
        with ThreadPoolExecutor(max_workers=4) as pool:
            list(pool.map(lambda _: self.db.save_job_record(self.job()), range(8)))
            reports = list(pool.map(lambda _: self.db.export_job_reports(), range(3)))
        self.assertEqual(len(self.db.get_reviewed_jobs()), 1)
        self.assertEqual(len({r["reviewed"] for r in reports}), 3)
        workbook = load_workbook(database_module.EXPORTS_DIR / "reviewed_jobs_latest.xlsx")
        self.assertEqual(workbook.active.max_row, 2)
        workbook.close()

    def test_spreadsheet_formula_text_stays_text(self):
        job = self.job(notes='=HYPERLINK("https://example.test", "click")')
        self.db.save_job_record(job)
        workbook = load_workbook(self.db.export_job_reports()["reviewed"])
        headers = {c.value: c.column for c in workbook.active[1]}
        cell = workbook.active.cell(2, headers["Notes"])
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, job.notes)
        workbook.close()

    def test_locked_latest_workbook_keeps_timestamped_report(self):
        self.db.save_job_record(self.job())
        real_replace = os.replace
        def replace(source, destination):
            if str(destination).endswith("_latest.xlsx"):
                raise PermissionError("Workbook is open")
            real_replace(source, destination)
        with patch("src.database.os.replace", side_effect=replace):
            reports = self.db.export_job_reports()
        self.assertTrue(all(path.exists() for path in reports.values()))


class ProfilePersistenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "profile.json"
        self.patcher = patch.object(config_module, "PROFILE_FILE", self.path)
        self.patcher.start()
        self.addCleanup(self.patcher.stop)

    def test_corrupt_profile_is_preserved(self):
        self.path.write_text('{broken JSON', encoding="utf-8")
        with self.assertRaisesRegex(RuntimeError, "file was preserved"):
            ConfigManager()
        self.assertEqual(self.path.read_text(encoding="utf-8"), '{broken JSON')

    def test_profile_restart_and_failed_write_rollback(self):
        manager = ConfigManager()
        manager.update_profile({"personal_info": {"full_name": "Fixture Candidate"}})
        before = self.path.read_bytes()
        with patch("src.config.os.replace", side_effect=PermissionError("locked")):
            with self.assertRaises(PermissionError):
                manager.update_profile({"personal_info": {"full_name": "Changed"}})
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(manager.profile["personal_info"]["full_name"], "Fixture Candidate")
        reopened = ConfigManager()
        self.assertEqual(reopened.profile["candidate_profile"]["personal_information"]["full_name"], "Fixture Candidate")

    def test_invalid_candidate_does_not_mutate_profile(self):
        manager = ConfigManager()
        before = copy.deepcopy(manager.profile)
        with self.assertRaises((ValueError, TypeError)):
            manager.update_profile({"candidate_profile": {"skills": 42}})
        self.assertEqual(manager.profile, before)


class RecoveryTests(unittest.IsolatedAsyncioTestCase):
    async def test_approval_timeout_expires_review_and_removes_waiter(self):
        from src.review import review_coordinator
        scraper = LinkedInScraper(MagicMock())
        with patch("src.scraper.db_manager.queue_application_review", return_value=999), patch("src.scraper.db_manager.resolve_application_review") as resolve, patch("src.scraper.asyncio.wait_for", new=AsyncMock(side_effect=asyncio.TimeoutError)):
            result = await scraper._await_user_approval(
                {"job_title": "Analyst", "company": "Fixture", "job_url": "https://example.test/job"},
                80, {"approval_timeout_minutes": 1}, AsyncMock())
        self.assertEqual(result["decision"], "expired")
        self.assertFalse(review_coordinator.is_active(999))
        self.assertEqual(resolve.call_args.args[1], "expired")

    async def test_failed_job_and_diagnostics_do_not_kill_next_job(self):
        scraper = LinkedInScraper(MagicMock())
        scraper._review_job = AsyncMock(side_effect=[RuntimeError("navigation failed"), {"status": "skipped"}])
        scraper._capture_failure_artifacts = AsyncMock(side_effect=PermissionError("disk locked"))
        scraper._dismiss_open_easy_apply_modal = AsyncMock()
        scraper._close_external_application = AsyncMock()
        totals = dict(applied=0, flagged=0, skipped=0, failed=0)
        jobs = [JobRecord(job_url=f"https://www.linkedin.com/jobs/view/{n}/") for n in (11, 12)]
        with patch("src.scraper.db_manager.is_job_known", return_value=False), patch("src.scraper.db_manager.save_job_record") as save:
            paused = await scraper._process_jobs(jobs, totals, 5, 0, ["Analyst"], {}, AsyncMock(), None)
        self.assertFalse(paused)
        self.assertEqual(totals["failed"], 1)
        self.assertEqual(totals["skipped"], 1)
        self.assertEqual(scraper._review_job.await_count, 2)
        self.assertEqual(save.call_args.args[0].status, "FAILED")

    async def test_browser_cleanup_continues_after_context_close_failure(self):
        manager = BrowserManager()
        manager.context = MagicMock(close=AsyncMock(side_effect=RuntimeError("already gone")))
        manager.playwright = MagicMock(stop=AsyncMock())
        engine = manager.playwright
        manager._owns_context = True
        await manager.close()
        engine.stop.assert_awaited_once()
        self.assertIsNone(manager.context)
        self.assertIsNone(manager.playwright)
        await manager.close()

    async def test_malformed_provider_envelopes_fall_back(self):
        for response in ({"choices": []}, [], {"choices": [None]}):
            provider = LMStudioProvider(LLMSettings(enabled=True, provider="lmstudio",
                                        base_url="http://localhost:1234/v1", model="fixture"),
                                        transport=AsyncMock(return_value=response))
            result = await provider.interpret_question("Unknown question", {})
            self.assertFalse(result.ok)
            self.assertTrue(result.needs_user_input)
            self.assertTrue(result.error)

    async def test_wrong_types_and_nonfinite_llm_values_are_rejected(self):
        data = dict(canonical_intent="email", answer="fixture@example.test", confidence=0.99,
                    needs_user_input=False, notes="Fixture")
        for updates in ({"confidence": float("nan")}, {"needs_user_input": "false"}, {"answer": ["Yes"]}):
            response = {"choices": [{"message": {"content": json.dumps({**data, **updates})}}]}
            provider = LMStudioProvider(LLMSettings(True, "lmstudio", "http://localhost:1234/v1", "fixture"),
                                        transport=AsyncMock(return_value=response))
            result = await provider.interpret_question("Question", {})
            self.assertFalse(result.ok)
            self.assertEqual(result.error, "invalid_schema")

    def test_explicit_llm_disable_overrides_provider_environment(self):
        with patch.dict(os.environ, {"LLM_PROVIDER": "lmstudio", "LLM_ENABLED": "false"}):
            self.assertFalse(LLMSettings.from_profile({}).enabled)
