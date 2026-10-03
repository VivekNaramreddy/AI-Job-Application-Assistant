import tempfile
import unittest
from pathlib import Path

import src.database as database_module
from openpyxl import load_workbook

from src.models import JobRecord


class ProcessedJobsDatabaseTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_file = database_module.DB_FILE
        self.original_exports_dir = database_module.EXPORTS_DIR
        base = Path(self.temp_dir.name)
        database_module.DB_FILE = base / "jobs.db"
        database_module.EXPORTS_DIR = base / "exports"
        self.database = database_module.DatabaseManager()

    def tearDown(self):
        database_module.DB_FILE = self.original_db_file
        database_module.EXPORTS_DIR = self.original_exports_dir
        self.temp_dir.cleanup()

    def test_rejection_details_are_atomic_durable_and_historical(self):
        job = JobRecord(job_url="https://www.linkedin.com/jobs/view/123999/",
                        status="SAVED_FOR_REVIEW", source_mode="search_save")
        self.database.save_job_record(job)
        for reason in (None, "", " ", "invalid"):
            with self.assertRaisesRegex(ValueError, "rejection reason"):
                self.database.set_manual_job_status(job.job_url, "Rejected", reason)
            self.assertEqual(self.database.get_dashboard_jobs()[0]["job_status"], "Pending")
        self.database.set_manual_job_status(job.job_url, "Rejected", "position_closed", "No longer accepting applications.")
        restarted = database_module.DatabaseManager()
        row = restarted.get_dashboard_jobs()[0]
        self.assertEqual(row["rejection_reason"], "position_closed")
        self.assertEqual(row["rejection_feedback"], "No longer accepting applications.")
        restarted.set_manual_job_status(job.job_url, "Rejected", "other", "")
        for status in ("Pending", "Applied"):
            restarted.set_manual_job_status(job.job_url, status)
            row = database_module.DatabaseManager().get_dashboard_jobs()[0]
            self.assertEqual(row["job_status"], status)
            self.assertEqual(row["rejection_reason"], "other")
            self.assertEqual(row["rejection_feedback"], "")
            self.assertEqual(row["source_mode"], "search_save")
            self.assertEqual(row["original_bucket"], "review")
        self.assertEqual(len(restarted.get_dashboard_jobs()), 1)

    def test_terminal_skip_is_known(self):
        url = "https://www.linkedin.com/jobs/view/1234567890/"
        self.database.log_processed_job(
            "Marketing Manager", "Example", url, "closed", "expired"
        )
        self.assertTrue(self.database.is_job_known(url))

    def test_delete_rechecks_filters_and_removes_auxiliary_rows(self):
        for index, (source, method, status) in enumerate((
            ("search_save", "external", "SAVED_FOR_REVIEW"),
            ("autopilot", "external", "REVIEWED"),
            ("search_save", "easy_apply", "SAVED_FOR_REVIEW"),
            ("search_save", "external", "SKIPPED"))):
            self.database.save_job_record(JobRecord(job_url=f"https://www.linkedin.com/jobs/view/{12000+index}/",
                status=status, source_mode=source, application_method=method))
        target = "https://www.linkedin.com/jobs/view/12000/"
        self.database.save_search_issue(JobRecord(job_url=target), "test", "Analyst", "Duplicate")
        with self.assertRaises(ValueError):
            self.database.delete_dashboard_jobs([target, target.replace("12000", "12001")], "review", "search_save", "external")
        self.assertEqual(len(self.database.get_dashboard_jobs()), 4)
        deleted = self.database.delete_dashboard_jobs([target + "?trackingId=1"], "review", "search_save", "external")
        self.assertEqual(deleted, [target])
        self.assertEqual(len(database_module.DatabaseManager().get_dashboard_jobs()), 3)
        self.assertEqual(self.database.get_search_issues(), [])

    def test_delete_protects_applied_history_even_after_status_changes(self):
        url = "https://www.linkedin.com/jobs/view/12005/"
        self.database.save_job_record(JobRecord(job_url=url, status="APPLIED"))
        for status, bucket in (("Applied", "applied"), ("Rejected", "rejected"), ("Pending", "review")):
            self.database.set_manual_job_status(url, status, "other")
            with self.database.get_connection() as conn:
                history = [dict(row) for row in conn.execute("SELECT * FROM applied_jobs")]
            with self.assertRaises(ValueError):
                self.database.delete_dashboard_jobs([url], bucket)
            with self.database.get_connection() as conn:
                self.assertEqual([dict(row) for row in conn.execute("SELECT * FROM applied_jobs")], history)
            self.assertTrue(self.database.get_dashboard_jobs()[0]["deletion_protected"])

    def test_delete_protects_live_review(self):
        url = "https://www.linkedin.com/jobs/view/12006/"
        self.database.queue_application_review("Analyst", "Example", url)
        with self.assertRaises(ValueError):
            self.database.delete_dashboard_jobs([url], "review")
        self.assertEqual(len(self.database.get_application_reviews()), 1)

    def test_application_metadata_survives_status_edits_and_restart(self):
        job = JobRecord(job_url="https://www.linkedin.com/jobs/view/123888/", source_mode="search_save",
                        application_method="external", application_url="https://ats.example.test/apply/1",
                        external_apply=True, status="SAVED_FOR_REVIEW")
        self.database.save_job_record(job)
        self.database.save_job_record(job)
        for status in ("Applied", "Rejected", "Pending"):
            self.database.set_manual_job_status(job.job_url, status, "other", "Custom rejection")
            rows = database_module.DatabaseManager().get_dashboard_jobs()
            self.assertEqual(len(rows), 1)
            row = rows[0]
            self.assertEqual(row["job_status"], status)
            self.assertEqual(row["source_mode"], "search_save")
            self.assertEqual(row["application_method"], "external")
            self.assertEqual(row["application_stage"], "not_started")
            self.assertEqual(row["external_apply_url"], job.application_url)
            self.assertEqual(row["original_bucket"], "review")

    def test_structured_external_failures(self):
        for reason, expected in (("captcha", "captcha"), ("login_required", "login_required"),
                                 ("unanswered_required_questions", "user_input_required"),
                                 ("application_closed", "application_closed"),
                                 ("form_not_found", "unsupported_ats")):
            job = JobRecord(job_url="https://www.linkedin.com/jobs/view/123777/", status="NEEDS_USER_INPUT",
                            application_method="external", ats_type="Generic HTML", failure_reason=reason)
            self.database.save_job_record(job)
            self.assertEqual(self.database.get_dashboard_jobs()[0]["automation_failure_reason"], expected)

    def test_application_metadata_migrates_existing_records(self):
        job = JobRecord(job_url="https://www.linkedin.com/jobs/view/123666/", status="SAVED_FOR_REVIEW",
                        source_mode="search_save", external_apply=True, application_url="https://ats.example.test/apply")
        self.database.save_job_record(job)
        with self.database.get_connection() as conn:
            for column in ("application_method", "application_stage", "automation_failure_reason"):
                conn.execute(f"ALTER TABLE processed_jobs DROP COLUMN {column}")
            conn.commit()
        row = database_module.DatabaseManager().get_dashboard_jobs()[0]
        self.assertEqual(row["application_method"], "external")
        self.assertEqual(row["application_stage"], "not_started")
        self.assertEqual(row["source_mode"], "search_save")

    def test_dashboard_status_moves_preserve_origin_and_metadata(self):
        job = JobRecord(job_url="https://www.linkedin.com/jobs/view/998877/",
                        title="Analyst", company="Example", status="REVIEWED",
                        description="Full description", notes="Match explanation", match_score=82)
        self.database.save_job_record(job)
        job.status = "SKIPPED"
        job.failure_reason = "Automatic decision"
        self.database.save_job_record(job)
        self.database.save_search_issue(job, "later-search", "Analyst", "Duplicate")
        for status, bucket in ((None, "skipped"), ("Applied", "applied"),
                               ("Rejected", "rejected"), ("Pending", "review")):
            if status:
                self.database.set_manual_job_status(job.job_url, status, "position_closed")
            rows = database_module.DatabaseManager().get_dashboard_jobs()
            self.assertEqual(len(rows), 1)
            row = rows[0]
            self.assertEqual(row["bucket"], bucket)
            self.assertEqual(row["original_bucket"], "skipped")
            self.assertEqual(row["source_mode"], "autopilot")
            self.assertEqual(row["id"], job.job_id)
            self.assertEqual(row["description"], "Full description")
            self.assertEqual(row["notes"], "Match explanation")
            self.assertEqual(row["match_score"], 82)
            self.assertEqual(row["failure_reason"], "Automatic decision")

    def test_dashboard_legacy_attempts_and_migration(self):
        job = JobRecord(job_url="https://www.linkedin.com/jobs/view/998878/",
                        title="Unreadable", description="Partial description", match_score=42)
        self.database.save_search_issue(job, "first", "Analyst", "Extraction Failed", "Timeout")
        self.database.save_search_issue(job, "second", "Analyst", "Extraction Failed", "Timeout")
        self.assertEqual(len(self.database.get_dashboard_jobs()), 1)
        self.database.set_manual_job_status(job.job_url, "Applied")
        row = self.database.get_dashboard_jobs()[0]
        self.assertEqual((row["source_mode"], row["original_bucket"], row["bucket"]),
                         ("search_save", "failed", "applied"))
        self.assertEqual(row["match_score"], 42)
        self.assertEqual(row["error"], "Timeout")
        with self.database.get_connection() as conn:
            conn.execute("UPDATE processed_jobs SET source_mode = NULL, original_bucket = NULL")
            conn.commit()
        restarted = database_module.DatabaseManager()
        self.assertEqual(restarted.get_dashboard_jobs()[0]["original_bucket"], "failed")
        self.assertEqual(restarted.get_dashboard_jobs()[0]["source_mode"], "unknown")

    def test_dashboard_normalized_status_ignores_original_classification(self):
        for index, (status, expected) in enumerate((
            (" applied ", "Applied"), ("rejected", "Rejected"),
            ("POSSIBLE_MATCH", "Pending"), ("rejected_by_preferences", "Skipped"),
            ("FAILED", "Failed"),
        )):
            job = JobRecord(job_url=f"https://www.linkedin.com/jobs/view/{990000 + index}/",
                            status=status, original_bucket="possible_match")
            self.database.save_job_record(job)
            row = next(j for j in self.database.get_dashboard_jobs() if j['job_url'] == job.job_url)
            self.assertEqual(row['job_status'], expected)
            self.assertEqual(row['original_bucket'], 'possible_match')

    def test_legacy_sources_require_evidence_and_do_not_change_on_status_update(self):
        for index, status in enumerate(('FAILED', 'SAVED_FOR_REVIEW')):
            self.database.save_job_record(JobRecord(
                job_url=f'https://www.linkedin.com/jobs/view/{991000 + index}/', status=status))
        with self.database.get_connection() as conn:
            conn.execute('UPDATE processed_jobs SET source_mode = NULL')
            conn.commit()
        self.database = database_module.DatabaseManager()
        for row in self.database.get_dashboard_jobs():
            expected = 'search_save' if row['status'] == 'SAVED_FOR_REVIEW' else 'unknown'
            self.assertEqual(row['source_mode'], expected)
            self.database.set_manual_job_status(row['job_url'], 'Applied')
            updated = next(j for j in self.database.get_dashboard_jobs() if j['job_url'] == row['job_url'])
            self.assertEqual(updated['source_mode'], expected)

    def test_manual_status_reuses_history_and_survives_restart(self):
        for status in ("Applied", "Rejected"):
            url = f"https://www.linkedin.com/jobs/view/{111111 if status == 'Applied' else 222222}/"
            job = JobRecord(job_url=url, title="Analyst", company="Example", status="POSSIBLE_MATCH")
            self.database.save_job_record(job)
            self.database.save_search_issue(job, "test", "Analyst", "Duplicate")
            result = self.database.set_manual_job_status(url + "?trackingId=test", status, "position_closed")
            self.assertEqual(result["tracking_status"], status)
            self.assertTrue(result["status_changed_at"])
            restarted = database_module.DatabaseManager()
            restarted.recover_interrupted_applications()
            saved = next(row for row in restarted.get_saved_jobs() if row["job_url"] == url)
            issue = next(row for row in restarted.get_search_issues() if row["job_url"] == url)
            self.assertEqual(saved["tracking_status"], status)
            self.assertEqual(issue["tracking_status"], status)
            self.assertEqual(saved["status_changed_at"], result["status_changed_at"])
            self.assertTrue(restarted.is_job_known(url, retry_legacy_search_rejections=True))
            self.assertTrue(restarted.is_job_blocked(url))
            job.status = "FAILED"
            restarted.save_job_record(job)
            self.assertEqual(job.status, status.upper(), "Late worker callbacks must preserve manual choices")
            pending = restarted.set_manual_job_status(url, "Pending")
            self.assertEqual(pending["status"], "POSSIBLE_MATCH")
            self.assertFalse(restarted.is_job_blocked(url))

    def test_manual_failed_job_status_is_shared_and_pending_remains_retryable(self):
        job = JobRecord(job_url="https://www.linkedin.com/jobs/view/333333/", title="Unreadable")
        self.database.save_search_issue(job, "run1", "Analyst", "Extraction Failed")
        self.database.save_search_issue(job, "run2", "Analyst", "Extraction Failed")
        self.assertEqual(self.database.get_search_issues()[0]["tracking_status"], "Pending")
        self.database.set_manual_job_status(job.job_url, "Rejected", "position_closed")
        self.assertTrue(all(row["tracking_status"] == "Rejected" for row in self.database.get_search_issues()))
        self.assertTrue(self.database.is_job_known(job.job_url))
        self.database.set_manual_job_status(job.job_url, "Pending")
        self.assertFalse(self.database.is_job_known(job.job_url))
        self.assertEqual(len(self.database.get_search_issues()), 2)

    def test_manual_change_preserves_legacy_rejections_and_closed_reviews(self):
        job = JobRecord(job_url="https://www.linkedin.com/jobs/view/444444/", title="Analyst",
                        status="NOT_QUALIFIED", notes="Score source: deterministic profile-aware")
        self.database.save_job_record(job)
        review_id = self.database.queue_application_review("Analyst", "Example", job.job_url)
        self.database.set_manual_job_status(job.job_url, "Rejected", "position_closed")
        self.assertTrue(self.database.is_job_known(job.job_url, retry_legacy_search_rejections=True))
        reviews = self.database.get_application_reviews(include_manually_tracked=True)
        self.assertEqual(reviews[0]["id"], review_id)
        self.assertEqual(reviews[0]["tracking_status"], "Rejected")
        self.assertEqual(reviews[0]["status"], "rejected")

    def test_status_change_cannot_race_an_approved_submission(self):
        url = "https://www.linkedin.com/jobs/view/555555/"
        review_id = self.database.queue_application_review("Analyst", "Example", url)
        self.database.resolve_application_review(review_id, "approved")
        with self.assertRaisesRegex(ValueError, "Submission is already in progress"):
            self.database.set_manual_job_status(url, "Rejected", "position_closed")

    def test_transient_failure_can_be_retried(self):
        url = "https://www.linkedin.com/jobs/view/9876543210/"
        self.database.log_processed_job(
            "Marketing Manager", "Example", url, "failed", "navigation timeout"
        )
        self.assertFalse(self.database.is_job_known(url))

    def test_exports_reviewed_and_applied_jobs_to_separate_workbooks(self):
        reviewed_url = "https://www.linkedin.com/jobs/view/1111111111/"
        applied_url = "https://www.linkedin.com/jobs/view/2222222222/"
        self.database.log_processed_job(
            "Content Writer",
            "Reviewed Company",
            reviewed_url,
            "rejected_by_preferences",
            "missing required keyword",
        )
        self.database.log_processed_job(
            "Marketing Manager",
            "Applied Company",
            applied_url,
            "applied",
            "Application submission confirmed",
        )
        self.database.log_applied_job(
            "Marketing Manager",
            "Applied Company",
            "Remote",
            "Full-time",
            applied_url,
        )

        reports = self.database.export_job_reports()

        self.assertEqual(set(reports), {"reviewed", "applied"})
        self.assertTrue(reports["reviewed"].exists())
        self.assertTrue(reports["applied"].exists())
        self.assertTrue((database_module.EXPORTS_DIR / "reviewed_jobs_latest.xlsx").exists())
        self.assertTrue((database_module.EXPORTS_DIR / "applied_jobs_latest.xlsx").exists())

        reviewed = database_module.pd.read_excel(
            reports["reviewed"], sheet_name="Jobs Reviewed"
        )
        applied = database_module.pd.read_excel(
            reports["applied"], sheet_name="Jobs Applied"
        )
        self.assertEqual(len(reviewed), 2)
        self.assertEqual(len(applied), 1)
        self.assertEqual(applied.iloc[0]["Job Title"], "Marketing Manager")
        self.assertEqual(applied.iloc[0]["Job URL"], "View Job")

        workbook = load_workbook(reports["reviewed"])
        sheet = workbook["Jobs Reviewed"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        url_cells = [
            sheet.cell(row=row, column=headers["Job URL"])
            for row in range(2, sheet.max_row + 1)
        ]
        targets = {cell.hyperlink.target for cell in url_cells if cell.hyperlink}
        self.assertIn(reviewed_url, targets)
        self.assertIn(applied_url, targets)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertTrue(sheet.auto_filter.ref)
        self.assertTrue(sheet["A1"].font.bold)

    def test_structured_record_preserves_metadata_and_missing_optional_fields(self):
        job = JobRecord(
            external_job_id="4444444444",
            title="Business Analyst",
            company="ABC Corporation",
            location="New York, NY",
            job_url="https://www.linkedin.com/jobs/view/4444444444/",
            source="LinkedIn",
            status="REVIEWED",
            date_reviewed="2026-09-02T12:30:00-04:00",
        )
        self.database.save_job_record(job)
        missing = JobRecord(
            job_url="https://www.linkedin.com/jobs/view/5555555555/",
            source="LinkedIn",
            status="REVIEWED",
        )
        self.database.save_job_record(missing)

        reports = self.database.export_job_reports()
        workbook = load_workbook(reports["reviewed"])
        sheet = workbook["Jobs Reviewed"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        rows = {
            sheet.cell(row=row, column=headers["External Job ID"]).value: row
            for row in range(2, sheet.max_row + 1)
        }
        job_row = rows["4444444444"]
        self.assertEqual(
            sheet.cell(job_row, headers["Job Title"]).value, "Business Analyst"
        )
        self.assertEqual(
            sheet.cell(job_row, headers["Company"]).value, "ABC Corporation"
        )
        self.assertEqual(sheet.cell(job_row, headers["Status"]).value, "REVIEWED")
        self.assertEqual(
            sheet.cell(job_row, headers["Job URL"]).hyperlink.target, job.job_url
        )
        self.assertEqual(len(rows), 2)

    def test_external_progress_and_start_date_survive_restart(self):
        job = JobRecord(
            job_url="https://www.linkedin.com/jobs/view/7777777777/",
            external_job_id="7777777777",
            title="Revenue Analyst",
            company="Example Corp",
            location="Boston, MA",
            application_url="https://jobs.lever.co/example/abc?utm_source=linkedin",
            ats_type="Lever",
            status="APPLICATION_STARTED",
            date_application_started="2026-09-03T10:30:00-04:00",
        )
        self.database.save_job_record(job)
        restarted = database_module.DatabaseManager()
        row = restarted.get_reviewed_jobs()[0]
        self.assertEqual(row["application_url"], job.application_url)
        self.assertEqual(row["ats_type"], "Lever")
        self.assertEqual(
            row["date_application_started"], job.date_application_started
        )
        reports = restarted.export_job_reports()
        workbook = load_workbook(reports["reviewed"])
        sheet = workbook["Jobs Reviewed"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        self.assertIn("Date Application Started", headers)
        external_cell = sheet.cell(row=2, column=headers["Application URL"])
        self.assertEqual(external_cell.value, "Open Application")
        self.assertEqual(external_cell.hyperlink.target, job.application_url)

    def test_external_url_and_metadata_combination_detect_duplicates(self):
        original = JobRecord(
            job_url="https://www.linkedin.com/jobs/view/8888888888/",
            title="Data Analyst",
            company="Example Corp",
            location="Remote",
            application_url="https://jobs.lever.co/example/abc?utm_source=linkedin",
            status="READY_FOR_REVIEW",
        )
        self.database.save_job_record(original)
        self.assertTrue(
            self.database.is_application_url_known(
                "https://jobs.lever.co/example/abc/?ref=linkedin",
                "https://www.linkedin.com/jobs/view/9999999999/",
            )
        )
        self.assertTrue(
            self.database.is_job_known(
                "https://www.linkedin.com/jobs/view/9999999999/",
                title="Data Analyst",
                company="Example Corp",
                location="Remote",
            )
        )

    def test_external_job_id_prevents_duplicate_review(self):
        original = JobRecord(
            job_url="https://www.linkedin.com/jobs/view/6666666666/",
            external_job_id="6666666666",
            title="Data Analyst",
            status="REVIEWED",
            date_reviewed="2026-09-02T12:30:00-04:00",
        )
        self.database.save_job_record(original)

        self.assertTrue(
            self.database.is_job_known(
                "https://www.linkedin.com/jobs/view/data-analyst-6666666666/",
                "6666666666",
            )
        )

    def test_application_review_queue_tracks_human_decision(self):
        review_id = self.database.queue_application_review(
            job_title="Product Marketing Manager",
            company="Example Company",
            job_url="https://www.linkedin.com/jobs/view/3333333333/",
            location="Remote",
            job_type="Full-time",
            application_system="LinkedIn Easy Apply",
            match_score=87,
            required_fields_complete=True,
            resume_uploaded=True,
            resume_used="candidate_resume.pdf",
            application_url="https://boards.greenhouse.io/example/jobs/333",
            final_submit_detected=True,
            field_count=6,
            step_count=3,
        )

        pending = self.database.get_application_reviews()
        self.assertEqual(len(pending), 1)
        self.assertEqual(pending[0]["id"], review_id)
        self.assertEqual(pending[0]["status"], "awaiting_approval")
        self.assertEqual(pending[0]["match_score"], 87)
        self.assertEqual(pending[0]["required_fields_complete"], 1)
        self.assertEqual(pending[0]["resume_uploaded"], 1)
        self.assertEqual(pending[0]["resume_used"], "candidate_resume.pdf")
        self.assertEqual(
            pending[0]["application_url"],
            "https://boards.greenhouse.io/example/jobs/333",
        )
        self.assertEqual(pending[0]["final_submit_detected"], 1)
        self.assertEqual(pending[0]["step_count"], 3)

        self.assertTrue(
            self.database.resolve_application_review(
                review_id, "rejected", "Rejected in dashboard"
            )
        )
        self.assertEqual(self.database.get_application_reviews(), [])
        history = self.database.get_application_reviews(status=None)
        self.assertEqual(history[0]["status"], "rejected")


if __name__ == "__main__":
    unittest.main()
