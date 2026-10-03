"""Offline system tests: real dashboard/API/Chromium, intercepted site fixtures.

No LinkedIn/ATS account or live application is contacted or submitted.
"""

import asyncio
import copy
import socket
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

import uvicorn
from playwright.async_api import async_playwright
from openpyxl import load_workbook

import src.database as database_module
import src.main as main_module
from src.browser import BrowserManager
from src.config import config_manager
from src.models import JobRecord
from src.form_filler import FormFiller
from src.ats.base import GenericHTMLAdapter
from src.ats.lever import LeverAdapter
from src.external_navigation import ExternalApplicationLauncher


class SystemBrowserTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        folder = Path(self.temp.name)
        self.patcher = patch.multiple(database_module, DB_FILE=folder / "jobs.db", EXPORTS_DIR=folder / "exports")
        self.patcher.start()
        self.addCleanup(self.patcher.stop)
        self.db = database_module.DatabaseManager()
        original = copy.deepcopy(config_manager.profile)
        self.addCleanup(setattr, config_manager, "profile", original)
        resume = folder / "candidate.pdf"
        resume.write_bytes(b"%PDF-1.7\nFixture resume")
        config_manager.profile = {
            "personal_info": {"email": "fixture@example.test", "phone": "5550100",
                              "resume_path": str(resume)},
            "work_preferences": {"target_roles": ["Analyst"], "target_locations": ["Remote"],
                                 "max_search_pages_per_query": 1, "max_applications_per_run": 1,
                                 "easy_apply_only": True, "approval_timeout_minutes": 1},
        }
        self.playwright = await async_playwright().start()
        self.addAsyncCleanup(self.playwright.stop)
        self.ui_browser = await self.playwright.chromium.launch(headless=True)
        self.addAsyncCleanup(self.ui_browser.close)
        self.manager = BrowserManager()
        manager_patch = patch.object(main_module, "browser_manager", self.manager)
        manager_patch.start()
        self.addCleanup(manager_patch.stop)
        profile_patch = patch("src.browser.USER_DATA_DIR", folder / "browser_profile")
        profile_patch.start()
        self.addCleanup(profile_patch.stop)
        self.external = False

        async def launch(log_callback=None):
            page = await self.manager.initialize(cdp_url="", headless=True, log_callback=log_callback)
            await page.context.route("**/*", self.site_route)
            return page
        launch_patch = patch.object(self.manager, "get_active_page", side_effect=launch)
        launch_patch.start()
        self.addCleanup(launch_patch.stop)
        sock = socket.socket()
        sock.bind(("127.0.0.1", 0))
        self.addCleanup(sock.close)
        self.base_url = f"http://127.0.0.1:{sock.getsockname()[1]}"
        self.server = uvicorn.Server(uvicorn.Config(main_module.app, log_level="error", access_log=False))
        self.server_task = asyncio.create_task(self.server.serve(sockets=[sock]))
        self.addAsyncCleanup(self.stop_server)
        async with asyncio.timeout(15):
            while not self.server.started:
                if self.server_task.done():
                    self.server_task.result()
                await asyncio.sleep(0.02)
        self.ui = await self.ui_browser.new_page()
        await self.ui.route("https://**/*", lambda route: route.abort())
        self.errors = []
        self.ui.on("pageerror", lambda error: self.errors.append(str(error)))
        self.ui.on("dialog", lambda dialog: dialog.accept())
        await self.ui.goto(self.base_url)
        self.assertEqual(self.errors, [], "Dashboard initialization must not throw")

    async def stop_server(self):
        self.server.should_exit = True
        await asyncio.wait_for(self.server_task, timeout=15)

    async def site_route(self, route):
        url = route.request.url
        form = """
          <form id="application-form">
            <label for="email">Email *</label><input id="email" type="email" required>
            <label for="phone">Phone *</label><input id="phone" type="tel" required>
            <label for="resume">Resume *</label><input id="resume" type="file" required>
            <button id="next" type="button" aria-label="Continue to next step">Next</button>
          </form>
          <script>
          window.submitClicks = 0;
          document.querySelector('#next').onclick = () => {
            window.capturedEmail = document.querySelector('#email').value;
            window.capturedResume = document.querySelector('#resume').value;
            document.querySelector('#application-form').innerHTML =
              '<h3>Review your application</h3><button type="button" id="submit" aria-label="Submit application">Submit application</button>';
            document.querySelector('#submit').onclick = () => {
              window.submitClicks++;
              document.querySelector('#application-form').innerHTML = '<h2>Your application was submitted</h2>';
            };
          };
          </script>
        """
        if "ats.example.test" in url:
            html = "<main>This ATS has no supported form</main>" if getattr(self, "unsupported_external", False) else form
        elif "/jobs/search/" in url:
            self.assertNotIn("f_AL=", url)
            html = '<main><li class="jobs-search-results__list-item"><a href="/jobs/view/5001/">Analyst</a></li></main>'
        elif "/jobs/view/" in url:
            apply = ('<a class="jobs-apply-button" href="https://ats.example.test/apply/5001">Apply</a>'
                     if self.external else '<button class="jobs-apply-button" aria-label="Easy Apply" onclick="document.querySelector(\'.jobs-easy-apply-modal\').style.display=\'block\'">Easy Apply</button>')
            html = f"""<main>
              <h1 class="job-details-jobs-unified-top-card__job-title">Analyst</h1>
              <div class="job-details-jobs-unified-top-card__company-name"><a>Fixture Company</a></div>
              <div class="job-details-jobs-unified-top-card__primary-description-container"><span>Remote</span></div>
              <article class="jobs-description"><div class="jobs-description__content">Analyze requirements. Full-time.</div></article>
              {apply}
              </main>"""
            if not self.external:
                html += '<div class="jobs-easy-apply-modal" role="dialog" aria-label="Easy Apply" style="display:none">' + form + '</div>'
        else:
            html = "<main>Fixture LinkedIn Jobs</main>"
        await route.fulfill(status=200, content_type="text/html", body=html)

    async def workflow(self, external, approved=True):
        self.external = external
        async with self.ui.expect_response("**/api/autopilot/start") as started:
            await self.ui.locator("#btn-start").click()
        self.assertEqual((await started.value).status, 200)
        await self.ui.locator('[data-tab="tab-reviews"]').click()
        async with asyncio.timeout(25):
            while not self.db.get_application_reviews():
                await asyncio.sleep(0.05)
        await self.ui.locator("#btn-refresh-reviews").click()
        await self.ui.get_by_role("button", name="Approve & Submit").wait_for()
        job = self.db.get_reviewed_jobs()[0]
        self.assertEqual(job["status"], "READY_FOR_REVIEW")
        self.assertEqual(job["job_title"], "Analyst")
        self.assertEqual(job["company"], "Fixture Company")
        self.assertEqual(job["application_method"], "external" if external else "easy_apply")
        self.assertEqual(job["source_mode"], "autopilot")
        self.assertEqual(job["application_stage"], "ready_for_review")
        self.assertEqual(self.db.get_applied_jobs(), [])
        application_page = self.manager.context.pages[-1] if external else self.manager.page
        self.assertEqual(await application_page.evaluate("window.submitClicks"), 0)
        self.assertEqual(await application_page.evaluate("window.capturedEmail"), "fixture@example.test")
        self.assertIn("candidate.pdf", await application_page.evaluate("window.capturedResume"))
        # Real API/report generation while the browser is awaiting user input.
        reviewed_response = await self.ui.request.get(self.base_url + "/api/export/reviewed")
        self.assertTrue(reviewed_response.ok)
        self.assertEqual((await self.ui.request.post(self.base_url + "/api/autopilot/start")).status, 200)
        await self.ui.get_by_role("button", name="Approve & Submit" if approved else "Reject", exact=True).click()
        async with asyncio.timeout(20):
            while main_module.autopilot_running:
                await asyncio.sleep(0.05)
        self.assertEqual(self.db.get_reviewed_jobs()[0]["status"], "APPLIED" if approved else "SKIPPED")
        self.assertEqual(len(self.db.get_applied_jobs()), int(approved))
        self.assertEqual(self.db.get_application_reviews(None)[0]["status"], "submitted" if approved else "rejected")
        if external:
            self.assertTrue(application_page.is_closed())
        else:
            self.assertEqual(await application_page.evaluate("window.submitClicks"), int(approved))
        workbook = load_workbook(database_module.EXPORTS_DIR / "applied_jobs_latest.xlsx")
        self.assertEqual(workbook.active.max_row, 1 + int(approved))
        workbook.close()
        self.assertEqual(self.errors, [])

    async def test_dashboard_to_easy_apply_confirmation_and_export(self):
        await self.workflow(False)

    async def test_dashboard_to_external_confirmation_and_export(self):
        await self.workflow(True)

    async def test_unsupported_external_is_preserved_without_duplicates(self):
        self.external = True
        self.unsupported_external = True
        response = await self.ui.request.post(self.base_url + "/api/autopilot/start")
        self.assertEqual(response.status, 200)
        async with asyncio.timeout(30):
            while main_module.autopilot_running:
                await asyncio.sleep(0.05)
        rows = self.db.get_dashboard_jobs()
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["application_method"], "external")
        self.assertEqual(row["automation_failure_reason"], "unsupported_ats")
        self.assertEqual(row["application_stage"], "failed")
        self.assertEqual(row["application_url"], "https://ats.example.test/apply/5001")
        self.assertEqual(row["bucket"], "failed")
        self.assertEqual(row["source_mode"], "autopilot")
        self.assertEqual(self.db.get_applied_jobs(), [])
        self.assertEqual(database_module.DatabaseManager().get_dashboard_jobs(), rows)
        self.assertEqual(self.errors, [])

    async def test_rejected_review_never_submits(self):
        await self.workflow(False, approved=False)

    async def test_generic_done_and_thanks_do_not_confirm_submission(self):
        await self.ui.set_content('<form><p>Thank you for your interest</p><button>Done</button></form>')
        self.assertFalse(await FormFiller(self.ui)._wait_for_submission_confirmation(self.ui.locator("form")))
        self.assertFalse(await GenericHTMLAdapter(self.ui, AsyncMock()).submission_confirmed())

    async def test_ats_domain_in_query_does_not_select_adapter(self):
        self.assertFalse(await LeverAdapter.detect(self.ui, "https://example.test/job?redirect=jobs.lever.co"))

    async def test_failed_external_navigation_closes_its_new_page(self):
        await self.ui.set_content('<a href="https://ats.example.test/broken">Apply</a>')
        before = len(self.ui.context.pages)
        await self.ui.context.route("https://ats.example.test/**", lambda route: route.abort())
        with self.assertRaises(Exception):
            await ExternalApplicationLauncher.open(self.ui, self.ui.locator("a"))
        self.assertEqual(len(self.ui.context.pages), before)
        self.assertFalse(self.ui.is_closed())
