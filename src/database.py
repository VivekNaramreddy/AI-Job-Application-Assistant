"""
src/database.py
=============================================================================
Database Manager & Excel Exporter

PYTHON LESSON:
- `sqlite3` is Python's built-in SQL relational database module.
- SQL (Structured Query Language) uses commands like `CREATE TABLE`, `INSERT`, 
  `SELECT`, and `UPDATE` to manage structured data in tables.
- `pandas` is Python's premier data library. We use `df.to_excel()` to convert SQL 
  tables into clean Excel `.xlsx` files.
=============================================================================
"""

import sqlite3
import sys
from contextlib import contextmanager, nullcontext
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional
import json
import logging
import os
import tempfile
from threading import RLock
from src.urls import canonical_url

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.models import JobRecord, current_timestamp

logger = logging.getLogger(__name__)
_export_lock = RLock()

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

from src.paths import BASE_DIR, DATA_DIR, EXPORTS_DIR
DB_FILE = DATA_DIR / "jobs.db"

REPORT_COLUMNS = [
    "Job ID",
    "External Job ID",
    "Job Title",
    "Company",
    "Location",
    "Status",
    "Match Score",
    "Source",
    "Workplace Type",
    "Employment Type",
    "Experience Level",
    "Salary",
    "Date Posted",
    "Date Discovered",
    "Date Reviewed",
    "Date Application Started",
    "Date Applied",
    "Resume Used",
    "Job URL",
    "Application URL",
    "ATS",
    "Failure Reason",
    "Notes",
]


REJECTION_REASONS = {'no_longer_accepting': 'No longer accepting applications',
 'position_closed': 'Position closed / filled',
 'job_expired': 'Job expired',
 'not_eligible': 'Not eligible / work authorization',
 'location_not_suitable': 'Location not suitable',
 'experience_too_high': 'Experience requirement too high',
 'skills_mismatch': 'Skills mismatch',
 'education_mismatch': 'Education mismatch',
 'compensation_not_suitable': 'Salary/compensation not suitable',
 'work_arrangement_not_suitable': 'Work arrangement not suitable',
 'duplicate_job': 'Duplicate job',
 'already_applied_elsewhere': 'Already applied elsewhere',
 'description_changed': 'Job description changed',
 'company_recruiter_issue': 'Company/recruiter issue',
 'not_interested': 'Not interested',
 'other': 'Other'}


class DatabaseManager:
    """Manages SQLite database for job history, flagged jobs queue, and Excel exports."""

    def __init__(self):
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        self.init_db()

    @contextmanager
    def get_connection(self) -> Iterator[sqlite3.Connection]:
        """Yield a SQLite connection and always close it afterward."""
        conn = sqlite3.connect(DB_FILE, timeout=30)
        conn.row_factory = sqlite3.Row  # Access columns by name like a Python dictionary
        try:
            yield conn
        finally:
            conn.close()

    def init_db(self) -> None:
        """Creates database tables if they do not exist."""
        with self.get_connection() as conn:
            cursor = conn.cursor()

            # Separate search attempts preserve saved/application history and
            # allow unreadable jobs to be retried on a later run.
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS search_job_issues (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    run_id TEXT NOT NULL, job_url TEXT NOT NULL,
                    external_job_id TEXT, job_title TEXT, company TEXT, location TEXT,
                    description TEXT, search_query TEXT, match_score INTEGER,
                    occurred_at TEXT NOT NULL, reason TEXT NOT NULL, error TEXT,
                    UNIQUE(run_id, job_url)
                )
            """)
            self._ensure_columns(conn, "search_job_issues", {
                "source_mode": "TEXT DEFAULT 'search_save'",
                "original_bucket": "TEXT",
            })
            conn.execute("""UPDATE search_job_issues SET original_bucket = CASE
                WHEN reason IN ('Duplicate', 'Closed', 'Low Match') THEN 'skipped'
                ELSE 'failed' END WHERE original_bucket IS NULL""")
            
            # Table for applied jobs
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS applied_jobs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    job_title TEXT NOT NULL,
                    company TEXT NOT NULL,
                    location TEXT,
                    job_type TEXT,
                    job_url TEXT UNIQUE,
                    description TEXT,
                    status TEXT DEFAULT 'Applied',
                    date_applied DATETIME DEFAULT CURRENT_TIMESTAMP,
                    notes TEXT
                )
            """)

            # Table for flagged / odd jobs requiring manual review
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS flagged_jobs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    job_title TEXT NOT NULL,
                    company TEXT NOT NULL,
                    job_url TEXT UNIQUE,
                    reason_flagged TEXT,
                    unhandled_questions TEXT, -- JSON string of unhandled questions
                    status TEXT DEFAULT 'Needs Review',
                    date_flagged DATETIME DEFAULT CURRENT_TIMESTAMP
                )
            """)

            # Every terminal job outcome is recorded here so closed, unsuitable,
            # external-apply, and already-applied postings are not retried on a
            # later run. Transient failures remain retryable.
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS processed_jobs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    job_title TEXT NOT NULL,
                    company TEXT NOT NULL,
                    location TEXT,
                    workplace_type TEXT,
                    employment_type TEXT,
                    job_url TEXT UNIQUE,
                    status TEXT NOT NULL,
                    reason TEXT,
                    match_score INTEGER,
                    processed_at DATETIME DEFAULT CURRENT_TIMESTAMP
                )
            """)
            self._ensure_columns(
                conn,
                "processed_jobs",
                {
                    "external_job_id": "TEXT",
                    "application_url": "TEXT",
                    "source": "TEXT",
                    "description": "TEXT",
                    "salary": "TEXT",
                    "experience_level": "TEXT",
                    "posted_date": "TEXT",
                    "date_discovered": "TEXT",
                    "date_reviewed": "TEXT",
                    "date_application_started": "TEXT",
                    "date_applied": "TEXT",
                    "resume_used": "TEXT",
                    "failure_reason": "TEXT",
                    "notes": "TEXT",
                    "ats_type": "TEXT",
                    "easy_apply": "INTEGER",
                    "external_apply": "INTEGER",
                    "status_before_manual": "TEXT",
                    "source_mode": "TEXT",
                    "original_bucket": "TEXT",
                    "status_changed_at": "TEXT",
                    "rejection_reason": "TEXT",
                    "rejection_feedback": "TEXT",
                    "application_method": "TEXT",
                    "application_stage": "TEXT",
                    "automation_failure_reason": "TEXT",
                },
            )
            conn.execute("""UPDATE processed_jobs SET application_method = CASE
                WHEN easy_apply = 1 THEN 'easy_apply'
                WHEN external_apply = 1 OR COALESCE(application_url, '') != '' THEN 'external'
                ELSE 'unavailable' END WHERE COALESCE(application_method, '') = ''""")
            conn.execute("""UPDATE processed_jobs SET application_stage = CASE
                WHEN COALESCE(status_before_manual, status) IN ('SAVED_FOR_REVIEW', 'POSSIBLE_MATCH') THEN 'not_started'
                ELSE LOWER(COALESCE(status_before_manual, status)) END
                WHERE COALESCE(application_stage, '') = ''""")
            cursor.execute(
                """
                UPDATE processed_jobs
                SET date_reviewed = COALESCE(NULLIF(date_reviewed, ''), processed_at),
                    source = CASE
                        WHEN COALESCE(source, '') = '' AND job_url LIKE '%linkedin.com/%'
                        THEN 'LinkedIn'
                        ELSE source
                    END,
                    failure_reason = CASE
                        WHEN COALESCE(failure_reason, '') != '' THEN failure_reason
                        WHEN UPPER(status) IN ('APPLIED', 'REVIEWED', 'QUALIFIED') THEN ''
                        ELSE COALESCE(NULLIF(failure_reason, ''), reason)
                    END
                """
            )

            # Applications that have been filled but must not be submitted
            # until the user explicitly approves them in the dashboard.
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS application_reviews (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    job_title TEXT NOT NULL,
                    company TEXT NOT NULL,
                    location TEXT,
                    job_type TEXT,
                    job_url TEXT UNIQUE,
                    application_system TEXT,
                    match_score INTEGER,
                    required_fields_complete INTEGER,
                    resume_uploaded INTEGER,
                    resume_used TEXT,
                    application_url TEXT,
                    final_submit_detected INTEGER,
                    field_count INTEGER,
                    step_count INTEGER,
                    status TEXT DEFAULT 'awaiting_approval',
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    decided_at DATETIME,
                    decision_notes TEXT
                )
            """)
            self._ensure_columns(
                conn,
                "application_reviews",
                {
                    "required_fields_complete": "INTEGER",
                    "resume_uploaded": "INTEGER",
                    "resume_used": "TEXT",
                    "application_url": "TEXT",
                    "final_submit_detected": "INTEGER",
                    "field_count": "INTEGER",
                    "step_count": "INTEGER",
                },
            )

            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS answer_memory (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    original_question TEXT NOT NULL,
                    normalized_question TEXT NOT NULL UNIQUE,
                    canonical_intent TEXT,
                    answer TEXT NOT NULL,
                    answer_source TEXT NOT NULL,
                    confidence REAL NOT NULL,
                    user_approved INTEGER NOT NULL DEFAULT 0,
                    last_used TEXT,
                    use_count INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL
                )
                """
            )
            cursor.execute(
                "CREATE INDEX IF NOT EXISTS idx_answer_memory_intent "
                "ON answer_memory(canonical_intent, user_approved)"
            )
            # Phase 4 external submissions saved a confirmed timestamp only in
            # processed_jobs. Repair that report omission without promoting reviews.
            conn.execute(
                """INSERT OR IGNORE INTO applied_jobs
                (job_title, company, location, job_type, job_url, description, status, date_applied, notes)
                SELECT job_title, company, location, employment_type, job_url, description,
                       'APPLIED', date_applied, notes FROM processed_jobs
                WHERE UPPER(status) = 'APPLIED' AND COALESCE(date_applied, '') != ''"""
            )
            conn.execute("UPDATE applied_jobs SET status = 'APPLIED' WHERE UPPER(status) = 'APPLIED'")
            # Infer legacy provenance once. Later status updates must preserve it.
            for row in conn.execute("SELECT * FROM processed_jobs WHERE source_mode IS NULL OR original_bucket IS NULL").fetchall():
                legacy_status = row["status_before_manual"] or row["status"]
                conn.execute("UPDATE processed_jobs SET source_mode = COALESCE(source_mode, ?), original_bucket = COALESCE(original_bucket, ?) WHERE id = ?",
                             (self._legacy_source(conn, dict(row)), self._workflow_bucket(legacy_status), row["id"]))
            conn.commit()
            logger.info("SQLite database initialized at %s", DB_FILE)

    @staticmethod
    def _ensure_columns(
        conn: sqlite3.Connection, table: str, columns: Dict[str, str]
    ) -> None:
        """Add new nullable columns without rebuilding an existing user table."""
        existing = {
            row["name"] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()
        }
        for column, sql_type in columns.items():
            if column not in existing:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {sql_type}")

    def log_applied_job(
        self,
        job_title: str,
        company: str,
        location: str,
        job_type: str,
        job_url: str,
        description: str = "",
        status: str = "Applied",
        notes: str = ""
    ) -> bool:
        """Logs a successfully applied job into the database."""
        if status.upper() != "APPLIED":
            raise ValueError("Applied history only accepts confirmed APPLIED submissions.")
        job_url = self.canonical_url(job_url)
        with self.get_connection() as conn:
            cursor = conn.cursor()
            try:
                cursor.execute("""
                    INSERT INTO applied_jobs (job_title, company, location, job_type, job_url, description, status, notes)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (job_title, company, location, job_type, job_url, description, "APPLIED", notes))
                conn.commit()
                logger.info("Confirmed application saved: %s at %s", job_title, company)
                return True
            except sqlite3.IntegrityError:
                # Job URL already exists in database
                logger.info("Application already saved: %s", job_url)
                return False

    def log_flagged_job(
        self,
        job_title: str,
        company: str,
        job_url: str,
        reason: str,
        unhandled_questions: Optional[List[Dict[str, Any]]] = None
    ) -> bool:
        """Logs an odd / complex job into the flagged jobs queue for manual review."""
        job_url = self.canonical_url(job_url)
        questions_json = json.dumps(unhandled_questions or [])
        with self.get_connection() as conn:
            cursor = conn.cursor()
            try:
                cursor.execute("""
                    INSERT INTO flagged_jobs (job_title, company, job_url, reason_flagged, unhandled_questions, status)
                    VALUES (?, ?, ?, ?, ?, 'Needs Review')
                    ON CONFLICT(job_url) DO UPDATE SET
                        job_title = excluded.job_title, company = excluded.company,
                        reason_flagged = excluded.reason_flagged,
                        unhandled_questions = excluded.unhandled_questions,
                        status = 'Needs Review', date_flagged = CURRENT_TIMESTAMP
                """, (job_title, company, job_url, reason, questions_json))
                conn.commit()
                logger.info("Job flagged for manual review: %s at %s", job_title, company)
                return True
            except sqlite3.IntegrityError:
                logger.exception("Could not save questions for %s", job_url)
                raise

    def get_applied_jobs(self, limit: int = 100) -> List[Dict[str, Any]]:
        """Fetches list of applied jobs sorted by date."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM applied_jobs WHERE UPPER(status) = 'APPLIED' ORDER BY date_applied DESC LIMIT ?", (limit,))
            rows = cursor.fetchall()
            return [dict(row) for row in rows]

    def get_reviewed_jobs(self, limit: int = 100) -> List[Dict[str, Any]]:
        """Fetch jobs the worker inspected, including their final outcomes."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "SELECT * FROM processed_jobs ORDER BY processed_at DESC LIMIT ?",
                (limit,),
            )
            rows = cursor.fetchall()
            return [dict(row) for row in rows]

    def save_search_issue(self, job: JobRecord, run_id: str, search_query: str,
                          reason: str, error: str = "") -> None:
        with self.get_connection() as conn:
            if reason == "Duplicate":
                existing = conn.execute(
                    "SELECT * FROM processed_jobs WHERE job_url = ? OR (? != '' AND external_job_id = ?) LIMIT 1",
                    (self.canonical_url(job.job_url), job.external_job_id, job.external_job_id),
                ).fetchone()
                if existing:
                    for field, column in (("title", "job_title"), ("company", "company"),
                                          ("location", "location"), ("description", "description"),
                                          ("external_job_id", "external_job_id")):
                        if not getattr(job, field):
                            setattr(job, field, existing[column] or "")
                    job.match_score = existing["match_score"]
            conn.execute("""
                INSERT INTO search_job_issues
                (run_id, job_url, external_job_id, job_title, company, location,
                 description, search_query, match_score, occurred_at, reason, error, original_bucket)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(run_id, job_url) DO NOTHING
            """, (run_id, self.canonical_url(job.job_url), job.external_job_id,
                  job.title, job.company, job.location, job.description, search_query,
                  job.match_score, current_timestamp(), reason, error,
                  "skipped" if reason in {"Duplicate", "Closed", "Low Match"} else "failed"))
            conn.commit()

    @staticmethod
    def _legacy_source(conn, item) -> str:
        if item.get("source_mode"):
            return item["source_mode"]
        history = str(item.get("status_before_manual") or item.get("status") or "").strip().upper()
        if history in {"SAVED_FOR_REVIEW", "POSSIBLE_MATCH"}:
            return "search_save"
        if conn.execute("SELECT 1 FROM application_reviews WHERE job_url = ?", (item["job_url"],)).fetchone():
            return "autopilot"
        return "unknown"

    @staticmethod
    def _workflow_bucket(status: str) -> str:
        status = str(status).strip().upper()
        if status in {"APPLIED", "ALREADY_APPLIED"}:
            return "applied"
        if status == "REJECTED":
            return "rejected"
        if status in {"SKIPPED", "NOT_QUALIFIED", "CLOSED", "DUPLICATE", "REJECTED_BY_PREFERENCES", "NO_EASY_APPLY", "REJECTED_BY_USER"}:
            return "skipped"
        if status in {"FAILED", "NEEDS_USER_INPUT"}:
            return "failed"
        return "review"

    @classmethod
    def _normalized_job_status(cls, item) -> str:
        bucket = cls._workflow_bucket(item["status"])
        if item.get("status_changed_at") and bucket not in {"applied", "rejected"}:
            return "Pending"
        return {"review": "Pending", "applied": "Applied", "rejected": "Rejected",
                "skipped": "Skipped", "failed": "Failed"}[bucket]

    def get_dashboard_jobs(self, connection=None) -> List[Dict[str, Any]]:
        """One view per canonical job; keep auxiliary attempt and approval history."""
        with (nullcontext(connection) if connection is not None else self.get_connection()) as conn:
            jobs = {}
            protected = {self.canonical_url(row[0]) for row in conn.execute("SELECT job_url FROM applied_jobs")}
            protected.update(self.canonical_url(row[0]) for row in conn.execute(
                "SELECT job_url FROM application_reviews WHERE status IN ('awaiting_approval', 'approved', 'submitted')"))
            for table, mode in (
                ("search_job_issues", "search_save"),
                ("flagged_jobs", "unknown"),
                ("application_reviews", "autopilot"),
                ("applied_jobs", "unknown"),
                ("processed_jobs", "unknown"),
            ):
                for row in conn.execute(f"SELECT * FROM {table} ORDER BY id"):
                    item = dict(row)
                    url = self.canonical_url(item["job_url"])
                    prior = jobs.get(url, {})
                    if table == "search_job_issues":
                        item["status"] = "SKIPPED" if item["reason"] in {"Duplicate", "Closed", "Low Match"} else "FAILED"
                        item["original_bucket"] = prior.get("original_bucket") or item["original_bucket"]
                    elif table == "flagged_jobs":
                        if item["status"] != "Needs Review":
                            continue
                        item["flagged_id"] = item["id"]
                        item["status"] = "FAILED"
                    elif table == "application_reviews":
                        item["review"] = dict(row)
                        item["status"] = "READY_FOR_REVIEW" if item["status"] in {"awaiting_approval", "approved"} else "SKIPPED"
                    item["source_mode"] = (self._legacy_source(conn, item) if table == "processed_jobs"
                                           else item.get("source_mode") or prior.get("source_mode") or mode)
                    item["original_bucket"] = item.get("original_bucket") or prior.get("original_bucket") or self._workflow_bucket(item["status"])
                    jobs[url] = {**prior, **item, "job_url": url}
            result = []
            for item in jobs.values():
                # Compatibility names share the existing stored URL/provider fields.
                item["external_apply_url"] = item.get("application_url") or ""
                item["ats_provider"] = item.get("ats_type") or ""
                item["job_status"] = self._normalized_job_status(item)
                item["bucket"] = self._workflow_bucket(item["job_status"])
                item["tracking_status"] = item["job_status"]
                item["application_method"] = item.get("application_method") or (
                    "easy_apply" if item.get("easy_apply") else "external" if item.get("external_apply") or item.get("application_url") else "unavailable")
                item["deletion_protected"] = (item["bucket"] == "applied" or item["job_url"] in protected
                    or bool(item.get("date_applied")) or item.get("status_before_manual") in {"APPLIED", "ALREADY_APPLIED"}
                    or item.get("status") in {"APPLICATION_FOUND", "APPLICATION_STARTED", "FORM_ANALYZED", "FILLING"})
                questions = item.get("unhandled_questions") or []
                if isinstance(questions, str):
                    try:
                        questions = json.loads(questions)
                    except (ValueError, TypeError):
                        questions = []
                item["unhandled_questions"] = questions
                result.append(item)
            return sorted(result, key=lambda item: item.get("processed_at") or item.get("occurred_at") or item.get("created_at") or "", reverse=True)

    def delete_dashboard_jobs(self, job_urls, bucket, source_mode="all", application_method="all"):
        """Delete an explicit visible snapshot atomically; never touch applied history."""
        if bucket not in {"review", "rejected", "skipped", "failed"}:
            raise ValueError("Jobs cannot be deleted from this tab.")
        if source_mode not in {"all", "autopilot", "search_save"} or application_method not in {"all", "easy_apply", "external"}:
            raise ValueError("Invalid job filters.")
        urls = {self.canonical_url(url) for url in job_urls}
        if not urls or "" in urls:
            raise ValueError("Select at least one job to delete.")
        with self.get_connection() as conn:
            conn.execute("BEGIN IMMEDIATE")
            eligible = {job["job_url"] for job in self.get_dashboard_jobs(conn)
                        if job["bucket"] == bucket and not job["deletion_protected"]
                        and (source_mode == "all" or job["source_mode"] == source_mode)
                        and (application_method == "all" or job["application_method"] == application_method)}
            if not urls.issubset(eligible):
                raise ValueError("The list changed or contains protected history. Refresh before deleting.")
            # Remove supporting rows too, so a deleted card cannot reappear from legacy history.
            # Compare canonical URLs for older rows with tracking parameters.
            for table in ("processed_jobs", "search_job_issues", "flagged_jobs", "application_reviews"):
                ids = [(row["id"],) for row in conn.execute(f"SELECT id, job_url FROM {table}")
                       if self.canonical_url(row["job_url"]) in urls]
                conn.executemany(f"DELETE FROM {table} WHERE id = ?", ids)
            conn.commit()
        return sorted(urls)

    def get_search_issues(self) -> List[Dict[str, Any]]:
        with self.get_connection() as conn:
            return self._with_tracking(conn, conn.execute(
                "SELECT * FROM search_job_issues ORDER BY id DESC"
            ).fetchall())

    @staticmethod
    def _with_tracking(conn, rows) -> List[Dict[str, Any]]:
        states = {row["job_url"]: row for row in conn.execute(
            "SELECT job_url, status, status_changed_at FROM processed_jobs"
        )}
        for row in conn.execute("SELECT job_url, status FROM applied_jobs WHERE UPPER(status) IN ('APPLIED', 'REJECTED')"):
            states.setdefault(row["job_url"], {"status": row["status"], "status_changed_at": None})
        result = []
        for row in rows:
            item = dict(row)
            state = states.get(DatabaseManager.canonical_url(item["job_url"]))
            status = str(state["status"] if state else "").upper()
            item["tracking_status"] = ("Applied" if status in {"APPLIED", "ALREADY_APPLIED"}
                                       else "Rejected" if status == "REJECTED" else "Pending")
            item["status_changed_at"] = state["status_changed_at"] if state else None
            result.append(item)
        return result

    def set_manual_job_status(self, job_url: str, status: str,
                              rejection_reason: Optional[str] = None,
                              rejection_feedback: Optional[str] = None) -> Dict[str, Any]:
        """Update the existing job outcome while retaining its workflow/history."""
        if status not in {"Pending", "Applied", "Rejected"}:
            raise ValueError("Unsupported job status")
        if status == "Rejected" and rejection_reason not in REJECTION_REASONS:
            raise ValueError("Select a valid rejection reason.")
        job_url = self.canonical_url(job_url)
        with self.get_connection() as conn:
            conn.execute("BEGIN IMMEDIATE")
            if conn.execute("SELECT 1 FROM application_reviews WHERE job_url = ? AND status = 'approved'",
                            (job_url,)).fetchone():
                raise ValueError("Submission is already in progress. Wait for its outcome before changing status.")
            job = conn.execute("SELECT * FROM processed_jobs WHERE job_url = ?", (job_url,)).fetchone()
            if job is None:
                # Extraction failures may exist only as search attempts. Create
                # their shared history row now, without making Pending failures
                # ineligible for a later retry.
                source = conn.execute("SELECT * FROM search_job_issues WHERE job_url = ? ORDER BY id DESC LIMIT 1",
                                      (job_url,)).fetchone()
                origin = "SKIPPED" if source and source["reason"] in {"Duplicate", "Closed", "Low Match"} else "FAILED"
                if source is None:
                    source = conn.execute("SELECT * FROM application_reviews WHERE job_url = ?", (job_url,)).fetchone()
                    origin = "READY_FOR_REVIEW"
                if source is None:
                    source = conn.execute("SELECT * FROM applied_jobs WHERE job_url = ?", (job_url,)).fetchone()
                    origin = "APPLIED"
                if source is None:
                    source = conn.execute("SELECT * FROM flagged_jobs WHERE job_url = ?", (job_url,)).fetchone()
                    origin = "FAILED"
                if source is None:
                    raise KeyError("Job was not found in review or search history")
                source = dict(source)
                first_attempt = conn.execute(
                    "SELECT source_mode, original_bucket FROM search_job_issues WHERE job_url = ? ORDER BY id LIMIT 1",
                    (job_url,),
                ).fetchone()
                if first_attempt:
                    source["source_mode"] = first_attempt["source_mode"]
                    source["original_bucket"] = first_attempt["original_bucket"]
                conn.execute("""INSERT INTO processed_jobs
                    (job_url, job_title, company, location, description, external_job_id, status, source)
                    VALUES (?, ?, ?, ?, ?, ?, ?, 'LinkedIn')""",
                    (job_url, source.get("job_title", ""), source.get("company", ""),
                     source.get("location", ""), source.get("description", ""),
                     source.get("external_job_id", ""), origin))
                # Preserve all compatible metadata when promoting a legacy-only
                # record; auxiliary attempts and reviews remain untouched.
                columns = {row["name"] for row in conn.execute("PRAGMA table_info(processed_jobs)")}
                metadata = {key: value for key, value in source.items()
                            if key in columns and key not in {"id", "job_url", "status"}}
                if metadata:
                    conn.execute("UPDATE processed_jobs SET " + ", ".join(f"{key} = ?" for key in metadata) + " WHERE job_url = ?",
                                 (*metadata.values(), job_url))
                job = conn.execute("SELECT * FROM processed_jobs WHERE job_url = ?", (job_url,)).fetchone()
            conn.execute("UPDATE processed_jobs SET source_mode = COALESCE(source_mode, ?), original_bucket = COALESCE(NULLIF(original_bucket, ''), ?) WHERE job_url = ?",
                         (self._legacy_source(conn, dict(job)),
                          self._workflow_bucket(job["status"]), job_url))
            previous = job["status_before_manual"] or job["status"]
            target = status.upper()
            if status == "Pending":
                target = previous if previous not in {"APPLIED", "ALREADY_APPLIED", "REJECTED"} else "FAILED"
            timestamp = current_timestamp()
            conn.execute("""UPDATE processed_jobs SET status = ?, status_before_manual = ?,
                status_changed_at = ?, date_applied = CASE WHEN ? = 'APPLIED'
                    THEN COALESCE(NULLIF(date_applied, ''), ?) ELSE date_applied END
                WHERE job_url = ?""", (target, previous, timestamp, target, timestamp, job_url))
            if status == "Rejected":
                conn.execute("""UPDATE processed_jobs SET rejection_reason = ?, rejection_feedback = ?
                    WHERE job_url = ?""", (rejection_reason, (rejection_feedback or "").strip(), job_url))
            if status == "Applied":
                conn.execute("""INSERT INTO applied_jobs
                    (job_title, company, location, job_type, job_url, description, status, date_applied, notes)
                    VALUES (?, ?, ?, ?, ?, ?, 'APPLIED', ?, 'Manually marked Applied')
                    ON CONFLICT(job_url) DO UPDATE SET status = 'APPLIED'""",
                    (job["job_title"], job["company"], job["location"], job["employment_type"],
                     job_url, job["description"], job["date_applied"] or timestamp))
            else:
                conn.execute("UPDATE applied_jobs SET status = ? WHERE job_url = ?", (status.upper(), job_url))
            if status != "Pending":
                conn.execute("""UPDATE application_reviews SET status = 'rejected', decided_at = ?,
                    decision_notes = 'Closed by manual job status'
                    WHERE job_url = ? AND status = 'awaiting_approval'""", (timestamp, job_url))
            row = conn.execute("SELECT * FROM processed_jobs WHERE job_url = ?", (job_url,)).fetchone()
            conn.commit()
            return self._with_tracking(conn, [row])[0]

    def is_job_blocked(self, job_url: str) -> bool:
        """Terminal user/application outcomes must never enter a new application."""
        with self.get_connection() as conn:
            return conn.execute("""SELECT 1 FROM processed_jobs WHERE job_url = ?
                AND UPPER(status) IN ('APPLIED', 'ALREADY_APPLIED', 'REJECTED')
                UNION ALL SELECT 1 FROM applied_jobs WHERE job_url = ?
                AND UPPER(status) IN ('APPLIED', 'REJECTED') LIMIT 1""",
                (self.canonical_url(job_url), self.canonical_url(job_url))).fetchone() is not None

    def get_saved_jobs(self) -> List[Dict[str, Any]]:
        """Search-only matches awaiting later review, not prepared applications."""
        with self.get_connection() as conn:
            rows = conn.execute(
                """SELECT * FROM processed_jobs WHERE status IN ('SAVED_FOR_REVIEW', 'POSSIBLE_MATCH')
                OR status_before_manual IN ('SAVED_FOR_REVIEW', 'POSSIBLE_MATCH') ORDER BY id DESC"""
            ).fetchall()
            return self._with_tracking(conn, rows)

    def queue_application_review(
        self,
        job_title: str,
        company: str,
        job_url: str,
        location: str = "",
        job_type: str = "",
        application_system: str = "",
        match_score: Optional[int] = None,
        required_fields_complete: bool = False,
        resume_uploaded: bool = False,
        resume_used: str = "",
        application_url: str = "",
        final_submit_detected: bool = False,
        field_count: Optional[int] = None,
        step_count: Optional[int] = None,
    ) -> int:
        """Persist a filled application that is waiting for human approval."""
        job_url = self.canonical_url(job_url)
        with self.get_connection() as conn:
            conn.execute(
                """
                INSERT INTO application_reviews (
                    job_title, company, location, job_type, job_url,
                    application_system, match_score, required_fields_complete,
                    resume_uploaded, resume_used, application_url, final_submit_detected,
                    field_count, step_count, status
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'awaiting_approval')
                ON CONFLICT(job_url) DO UPDATE SET
                    job_title = excluded.job_title,
                    company = excluded.company,
                    location = excluded.location,
                    job_type = excluded.job_type,
                    application_system = excluded.application_system,
                    match_score = excluded.match_score,
                    required_fields_complete = excluded.required_fields_complete,
                    resume_uploaded = excluded.resume_uploaded,
                    resume_used = excluded.resume_used,
                    application_url = excluded.application_url,
                    final_submit_detected = excluded.final_submit_detected,
                    field_count = excluded.field_count,
                    step_count = excluded.step_count,
                    status = 'awaiting_approval',
                    created_at = CURRENT_TIMESTAMP,
                    decided_at = NULL,
                    decision_notes = NULL
                """,
                (
                    job_title,
                    company,
                    location,
                    job_type,
                    job_url,
                    application_system,
                    match_score,
                    int(required_fields_complete),
                    int(resume_uploaded),
                    resume_used,
                    application_url,
                    int(final_submit_detected),
                    field_count,
                    step_count,
                ),
            )
            row = conn.execute(
                "SELECT id FROM application_reviews WHERE job_url = ?", (job_url,)
            ).fetchone()
            conn.commit()
            return int(row["id"])

    def get_application_reviews(
        self, status: Optional[str] = "awaiting_approval", include_manually_tracked: bool = False
    ) -> List[Dict[str, Any]]:
        """Return application reviews, optionally filtered by workflow status."""
        with self.get_connection() as conn:
            if status is None:
                rows = conn.execute(
                    "SELECT * FROM application_reviews ORDER BY created_at DESC"
                ).fetchall()
            else:
                rows = conn.execute(
                    """
                    SELECT r.* FROM application_reviews r
                    LEFT JOIN processed_jobs p ON p.job_url = r.job_url
                    WHERE r.status = ? OR (? AND p.status_changed_at IS NOT NULL)
                    ORDER BY r.created_at DESC
                    """,
                    (status, include_manually_tracked),
                ).fetchall()
            return self._with_tracking(conn, rows)

    def resolve_application_review(
        self, review_id: int, status: str, notes: str = ""
    ) -> bool:
        """Record a review decision or the final submission outcome."""
        allowed = {
            "approved",
            "rejected",
            "expired",
            "submitted",
            "submission_failed",
        }
        if status not in allowed:
            raise ValueError(f"Unsupported application review status: {status}")
        predecessors = {
            "approved": ("awaiting_approval",),
            "rejected": ("awaiting_approval",),
            "expired": ("awaiting_approval", "approved"),
            "submitted": ("approved",),
            "submission_failed": ("awaiting_approval", "approved"),
        }[status]
        with self.get_connection() as conn:
            cursor = conn.execute(
                f"""
                UPDATE application_reviews
                SET status = ?, decided_at = CURRENT_TIMESTAMP, decision_notes = ?
                WHERE id = ? AND status IN ({','.join('?' for _ in predecessors)})
                """,
                (status, notes, review_id, *predecessors),
            )
            conn.commit()
            return cursor.rowcount > 0

    def is_job_known(
        self,
        job_url: str,
        external_job_id: str = "",
        title: str = "",
        company: str = "",
        location: str = "",
        retry_legacy_search_rejections: bool = False,
    ) -> bool:
        """Return True for a stable ID/URL or a high-confidence prior outcome."""
        job_url = self.canonical_url(job_url)
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                """
                SELECT 1 FROM applied_jobs WHERE job_url = ? AND UPPER(status) IN ('APPLIED', 'REJECTED')
                UNION ALL
                SELECT 1 FROM flagged_jobs WHERE job_url = ? AND status = 'Needs Review'
                UNION ALL
                SELECT 1 FROM processed_jobs
                WHERE (job_url = ? OR (? != '' AND external_job_id = ?))
                  AND UPPER(status) != 'FAILED'
                  AND NOT (? AND UPPER(status) = 'NOT_QUALIFIED'
                           AND COALESCE(notes, '') NOT LIKE '%Score source: weighted search v1%')
                LIMIT 1
                """,
                (job_url, job_url, job_url, external_job_id, external_job_id, retry_legacy_search_rejections),
            )
            if cursor.fetchone() is not None:
                return True
            if not all(value.strip() for value in (title, company, location)):
                return False
            cursor.execute(
                """
                SELECT 1 FROM processed_jobs
                WHERE LOWER(TRIM(job_title)) = LOWER(TRIM(?))
                  AND LOWER(TRIM(company)) = LOWER(TRIM(?))
                  AND LOWER(TRIM(location)) = LOWER(TRIM(?))
                  AND UPPER(status) IN (
                    'APPLICATION_STARTED', 'NEEDS_USER_INPUT', 'READY_FOR_REVIEW',
                    'APPLIED', 'ALREADY_APPLIED', 'REJECTED', 'DUPLICATE'
                  )
                LIMIT 1
                """,
                (title, company, location),
            )
            return cursor.fetchone() is not None

    def is_application_url_known(
        self, application_url: str, exclude_job_url: str = ""
    ) -> bool:
        """Compare canonical external URLs while allowing the current record."""
        canonical = self.canonical_url(application_url)
        if not canonical:
            return False
        with self.get_connection() as conn:
            rows = conn.execute(
                """
                SELECT job_url, application_url, status FROM processed_jobs
                WHERE COALESCE(application_url, '') != '' AND UPPER(status) != 'FAILED'
                """
            ).fetchall()
        return any(
            row["job_url"] != exclude_job_url
            and self.canonical_url(row["application_url"]) == canonical
            for row in rows
        )

    canonical_url = staticmethod(canonical_url)

    @staticmethod
    def _automation_failure_code(job: JobRecord) -> str:
        """Structured diagnostics, separate from matching and manual rejection reasons."""
        if job.source_mode == "search_save" or job.status not in {"FAILED", "NEEDS_USER_INPUT", "CLOSED"}:
            return ""
        reason = job.failure_reason.lower()
        if job.status == "CLOSED" or "application_closed" in reason:
            return "application_closed"
        if "captcha" in reason or "security" in reason:
            return "captcha"
        if "login" in reason or "sign-in" in reason:
            return "login_required"
        if "form_not_found" in reason or "no supported application form" in reason:
            return "unsupported_ats" if job.ats_type == "Generic HTML" else "unknown_form"
        if job.status == "NEEDS_USER_INPUT":
            return "user_input_required"
        return "unknown_form" if job.application_method == "external" else "automation_failed"

    def save_job_record(self, job: JobRecord) -> int:
        """Create or update one normalized job-history record."""
        job.job_url = self.canonical_url(job.job_url)
        if not job.job_url:
            raise ValueError("A job URL is required to persist history.")
        job.status = job.status.upper()
        job.application_method = job.application_method or (
            "easy_apply" if job.easy_apply else "external" if job.external_apply or job.application_url else "unavailable"
        )
        job.application_stage = ("not_started" if job.status in {"SAVED_FOR_REVIEW", "POSSIBLE_MATCH"}
                                 else job.status.lower())
        job.automation_failure_reason = self._automation_failure_code(job)
        if job.status == "APPLIED" and not job.date_applied:
            job.date_applied = current_timestamp()
        with self.get_connection() as conn:
            # A late callback/retry must never demote a confirmed submission.
            first_attempt = conn.execute(
                "SELECT source_mode, original_bucket FROM search_job_issues WHERE job_url = ? ORDER BY id LIMIT 1",
                (job.job_url,),
            ).fetchone()
            if first_attempt and not conn.execute("SELECT 1 FROM processed_jobs WHERE job_url = ?", (job.job_url,)).fetchone():
                job.source_mode = first_attempt["source_mode"] or "search_save"
                job.original_bucket = first_attempt["original_bucket"] or "failed"
            existing = conn.execute(
                """SELECT id, status FROM processed_jobs WHERE job_url = ? AND
                (UPPER(status) = 'APPLIED' OR (status_changed_at IS NOT NULL AND UPPER(status) = 'REJECTED'))""",
                (job.job_url,),
            ).fetchone()
            if existing and (job.status != "APPLIED" or existing["status"] == "REJECTED"):
                job.job_id = int(existing["id"])
                job.status = existing["status"]
                return job.job_id
            conn.execute(
                """
                INSERT INTO processed_jobs (
                    job_title, company, location, workplace_type, employment_type,
                    job_url, status, reason, match_score, external_job_id,
                    application_url, source, description, salary, experience_level,
                    posted_date, date_discovered, date_reviewed, date_applied,
                    date_application_started, resume_used, failure_reason, notes, ats_type, easy_apply,
                    external_apply, source_mode, original_bucket,
                    application_method, application_stage, automation_failure_reason
                )
                VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                )
                ON CONFLICT(job_url) DO UPDATE SET
                    job_title = excluded.job_title,
                    company = excluded.company,
                    location = excluded.location,
                    workplace_type = excluded.workplace_type,
                    employment_type = excluded.employment_type,
                    status = excluded.status,
                    reason = excluded.reason,
                    match_score = excluded.match_score,
                    external_job_id = excluded.external_job_id,
                    application_url = COALESCE(NULLIF(excluded.application_url, ''), processed_jobs.application_url),
                    source = excluded.source,
                    description = excluded.description,
                    salary = excluded.salary,
                    experience_level = excluded.experience_level,
                    posted_date = excluded.posted_date,
                    date_discovered = COALESCE(
                        NULLIF(processed_jobs.date_discovered, ''), excluded.date_discovered
                    ),
                    date_reviewed = excluded.date_reviewed,
                    date_application_started = excluded.date_application_started,
                    date_applied = excluded.date_applied,
                    resume_used = excluded.resume_used,
                    failure_reason = excluded.failure_reason,
                    notes = excluded.notes,
                    ats_type = excluded.ats_type,
                    easy_apply = excluded.easy_apply,
                    external_apply = excluded.external_apply,
                    application_method = excluded.application_method,
                    application_stage = excluded.application_stage,
                    automation_failure_reason = excluded.automation_failure_reason,
                    original_bucket = COALESCE(NULLIF(processed_jobs.original_bucket, ''), excluded.original_bucket),
                    processed_at = CURRENT_TIMESTAMP
                WHERE UPPER(processed_jobs.status) != 'APPLIED' OR excluded.status = 'APPLIED'
                """,
                (
                    job.title,
                    job.company,
                    job.location,
                    job.workplace_type,
                    job.employment_type,
                    job.job_url,
                    job.status.upper(),
                    job.failure_reason,
                    job.match_score,
                    job.external_job_id,
                    job.application_url,
                    job.source,
                    job.description,
                    job.salary,
                    job.experience_level,
                    job.posted_date,
                    job.date_discovered,
                    job.date_reviewed,
                    job.date_applied,
                    job.date_application_started,
                    job.resume_used,
                    job.failure_reason,
                    job.notes,
                    job.ats_type,
                    None if job.easy_apply is None else int(job.easy_apply),
                    None if job.external_apply is None else int(job.external_apply),
                    job.source_mode,
                    job.original_bucket or ("" if job.status in {
                        "DISCOVERED", "REVIEWED", "QUALIFIED", "APPLICATION_FOUND",
                        "APPLICATION_STARTED", "FORM_ANALYZED", "FILLING"
                    } else self._workflow_bucket(job.status)),
                    job.application_method,
                    job.application_stage,
                    job.automation_failure_reason,
                ),
            )
            row = conn.execute(
                "SELECT id FROM processed_jobs WHERE job_url = ?", (job.job_url,)
            ).fetchone()
            if job.status == "APPLIED":
                conn.execute(
                    """INSERT INTO applied_jobs
                    (job_title, company, location, job_type, job_url, description, status, date_applied, notes)
                    VALUES (?, ?, ?, ?, ?, ?, 'APPLIED', ?, ?)
                    ON CONFLICT(job_url) DO NOTHING""",
                    (job.title, job.company, job.location, job.employment_type,
                     job.job_url, job.description, job.date_applied or current_timestamp(), job.notes),
                )
            conn.commit()
            job.job_id = int(row["id"])
            return job.job_id

    def log_processed_job(
        self,
        job_title: str,
        company: str,
        job_url: str,
        status: str,
        reason: str = "",
        location: str = "",
        workplace_type: str = "",
        employment_type: str = "",
        match_score: Optional[int] = None,
    ) -> None:
        """Backward-compatible wrapper around structured job-history storage."""
        record = JobRecord(
            title=job_title,
            company=company,
            location=location,
            workplace_type=workplace_type,
            employment_type=employment_type,
            job_url=job_url,
            source="LinkedIn" if "linkedin.com/" in job_url else "",
            status=status.upper(),
            match_score=match_score,
            date_reviewed=current_timestamp(),
            failure_reason=reason,
        )
        self.save_job_record(record)

    def get_flagged_jobs(self, status: str = "Needs Review") -> List[Dict[str, Any]]:
        """Fetches jobs waiting for manual review."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM flagged_jobs WHERE status = ? ORDER BY date_flagged DESC", (status,))
            rows = cursor.fetchall()
            result = []
            for row in rows:
                r = dict(row)
                if r.get("unhandled_questions"):
                    try:
                        r["unhandled_questions"] = json.loads(r["unhandled_questions"])
                    except (json.JSONDecodeError, TypeError):
                        logger.warning("Invalid question JSON in flagged job %s", r["id"])
                        r["unhandled_questions"] = []
                result.append(r)
            return result

    def resolve_flagged_job(self, flagged_id: int) -> bool:
        """Marks a flagged job as resolved after user review."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("UPDATE flagged_jobs SET status = 'Resolved' WHERE id = ?", (flagged_id,))
            resolved = cursor.rowcount > 0
            conn.execute(
                """UPDATE processed_jobs SET status = 'FAILED', failure_reason = 'Answers updated; retry available'
                WHERE job_url = (SELECT job_url FROM flagged_jobs WHERE id = ?)
                  AND UPPER(status) IN ('NEEDS_USER_INPUT', 'FAILED')""",
                (flagged_id,),
            )
            conn.commit()
            return resolved

    def recover_interrupted_applications(self) -> None:
        """Run once at server startup, never when a report reader is created."""
        with self.get_connection() as conn:
            conn.execute(
                """UPDATE application_reviews SET status = 'expired',
                decided_at = CURRENT_TIMESTAMP,
                decision_notes = 'Server restarted; review the application again'
                WHERE status IN ('awaiting_approval', 'approved')"""
            )
            conn.execute(
                """UPDATE processed_jobs SET status = 'FAILED',
                failure_reason = 'Session interrupted; application was not confirmed. Verify on site before retrying.'
                WHERE UPPER(status) IN ('APPLICATION_FOUND', 'APPLICATION_STARTED',
                    'FORM_ANALYZED', 'FILLING', 'READY_FOR_REVIEW')"""
            )
            conn.commit()

    def export_job_reports(self) -> Dict[str, Path]:
        """Create separate Excel files for inspected jobs and confirmed applications."""
        with _export_lock:
            return self._export_job_reports()

    def _export_job_reports(self) -> Dict[str, Path]:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        report_paths = {
            "reviewed": EXPORTS_DIR / f"reviewed_jobs_{timestamp}.xlsx",
            "applied": EXPORTS_DIR / f"applied_jobs_{timestamp}.xlsx",
        }
        latest_paths = {
            "reviewed": EXPORTS_DIR / "reviewed_jobs_latest.xlsx",
            "applied": EXPORTS_DIR / "applied_jobs_latest.xlsx",
        }

        with self.get_connection() as conn:
            conn.execute("BEGIN")
            reviewed_jobs = self._read_reviewed_report(conn)
            applied_jobs = self._read_applied_report(conn)

        self._write_excel_report(
            reviewed_jobs,
            report_paths["reviewed"],
            latest_paths["reviewed"],
            "Jobs Reviewed",
        )
        self._write_excel_report(
            applied_jobs,
            report_paths["applied"],
            latest_paths["applied"],
            "Jobs Applied",
        )

        print(
            "Exported separate job reports to "
            f"{report_paths['reviewed']} and {report_paths['applied']}"
        )
        return report_paths

    @staticmethod
    def _read_reviewed_report(conn: sqlite3.Connection) -> pd.DataFrame:
        return pd.read_sql_query(
            """
            SELECT
                id AS 'Job ID',
                external_job_id AS 'External Job ID',
                job_title AS 'Job Title',
                company AS 'Company',
                location AS 'Location',
                UPPER(status) AS 'Status',
                match_score AS 'Match Score',
                source AS 'Source',
                workplace_type AS 'Workplace Type',
                employment_type AS 'Employment Type',
                experience_level AS 'Experience Level',
                salary AS 'Salary',
                posted_date AS 'Date Posted',
                date_discovered AS 'Date Discovered',
                COALESCE(NULLIF(date_reviewed, ''), processed_at) AS 'Date Reviewed',
                date_application_started AS 'Date Application Started',
                date_applied AS 'Date Applied',
                resume_used AS 'Resume Used',
                job_url AS 'Job URL',
                application_url AS 'Application URL',
                ats_type AS 'ATS',
                COALESCE(NULLIF(failure_reason, ''), reason) AS 'Failure Reason',
                notes AS 'Notes'
            FROM processed_jobs
            ORDER BY COALESCE(NULLIF(date_reviewed, ''), processed_at) DESC
            """,
            conn,
        ).reindex(columns=REPORT_COLUMNS)

    @staticmethod
    def _read_applied_report(conn: sqlite3.Connection) -> pd.DataFrame:
        """Preserve legacy applied rows while enriching them from structured history."""
        return pd.read_sql_query(
            """
            SELECT
                p.id AS 'Job ID',
                p.external_job_id AS 'External Job ID',
                COALESCE(NULLIF(p.job_title, ''), a.job_title) AS 'Job Title',
                COALESCE(NULLIF(p.company, ''), a.company) AS 'Company',
                COALESCE(NULLIF(p.location, ''), a.location) AS 'Location',
                'APPLIED' AS 'Status',
                p.match_score AS 'Match Score',
                COALESCE(NULLIF(p.source, ''), 'LinkedIn') AS 'Source',
                p.workplace_type AS 'Workplace Type',
                COALESCE(NULLIF(p.employment_type, ''), a.job_type) AS 'Employment Type',
                p.experience_level AS 'Experience Level',
                p.salary AS 'Salary',
                p.posted_date AS 'Date Posted',
                p.date_discovered AS 'Date Discovered',
                p.date_reviewed AS 'Date Reviewed',
                p.date_application_started AS 'Date Application Started',
                COALESCE(NULLIF(p.date_applied, ''), a.date_applied) AS 'Date Applied',
                p.resume_used AS 'Resume Used',
                a.job_url AS 'Job URL',
                p.application_url AS 'Application URL',
                p.ats_type AS 'ATS',
                '' AS 'Failure Reason',
                COALESCE(NULLIF(p.notes, ''), a.notes) AS 'Notes'
            FROM applied_jobs a
            LEFT JOIN processed_jobs p ON p.job_url = a.job_url
            WHERE UPPER(a.status) = 'APPLIED'
            ORDER BY COALESCE(NULLIF(p.date_applied, ''), a.date_applied) DESC
            """,
            conn,
        ).reindex(columns=REPORT_COLUMNS)

    @staticmethod
    def _write_excel_report(
        dataframe: pd.DataFrame,
        timestamped_path: Path,
        latest_path: Path,
        sheet_name: str,
    ) -> None:
        """Write readable timestamped/latest workbooks with clickable links."""
        dataframe = dataframe.copy()
        for column in (
            "Date Discovered", "Date Reviewed", "Date Application Started", "Date Applied"
        ):
            if column in dataframe:
                dataframe[column] = dataframe[column].map(
                    DatabaseManager._excel_datetime
                )
        for destination in (timestamped_path, latest_path):
            handle, name = tempfile.mkstemp(suffix=".xlsx", dir=destination.parent)
            os.close(handle)
            temporary = Path(name)
            try:
                with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
                    dataframe.to_excel(writer, sheet_name=sheet_name, index=False)
                    worksheet = writer.book[sheet_name]
                    # Treat all scraped text as text, including leading '='.
                    for row in worksheet.iter_rows(min_row=2):
                        for cell in row:
                            if cell.data_type == "f":
                                cell.data_type = "s"
                    DatabaseManager._format_excel_worksheet(worksheet)
                os.replace(temporary, destination)
            except PermissionError:
                if destination == timestamped_path:
                    raise
                logger.warning("Close %s in Excel to refresh the latest copy; timestamped report was saved.", destination)
            finally:
                temporary.unlink(missing_ok=True)

    @staticmethod
    def _excel_datetime(value):
        """Convert stored timestamps to Excel-safe local wall-clock datetimes."""
        if value is None or str(value).strip() == "":
            return None
        parsed = pd.to_datetime(value, errors="coerce")
        if pd.isna(parsed):
            return None
        if getattr(parsed, "tzinfo", None) is not None:
            parsed = parsed.tz_localize(None)
        return parsed.to_pydatetime()

    @staticmethod
    def _format_excel_worksheet(worksheet) -> None:
        """Apply lightweight report formatting without changing underlying data."""
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        headers = {cell.value: cell.column for cell in worksheet[1]}

        for label, display in (("Job URL", "View Job"), ("Application URL", "Open Application")):
            column_index = headers.get(label)
            if not column_index:
                continue
            for row in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=column_index)
                url = str(cell.value or "").strip()
                if url.startswith(("http://", "https://")):
                    cell.hyperlink = url
                    cell.value = display
                    cell.style = "Hyperlink"

        date_columns = {
            "Date Discovered", "Date Reviewed", "Date Application Started", "Date Applied"
        }
        wrap_columns = {"Failure Reason", "Notes"}
        for column_index in range(1, worksheet.max_column + 1):
            header = worksheet.cell(row=1, column=column_index).value
            max_length = len(str(header or ""))
            for row in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=column_index)
                max_length = max(max_length, len(str(cell.value or "")))
                if header in date_columns and cell.value is not None:
                    cell.number_format = "yyyy-mm-dd hh:mm"
                if header in wrap_columns:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            width_cap = 48 if header in wrap_columns else 32
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                width_cap, max(11, max_length + 2)
            )

    def export_to_excel(self) -> Path:
        """Backward-compatible applied-jobs export used by older callers."""
        return self.export_job_reports()["applied"]


db_manager = DatabaseManager()
